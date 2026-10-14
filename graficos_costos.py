"""Visualiza el costo estimado de alimento por kg de crecimiento bruto."""

import csv
import math
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
from openpyxl import load_workbook


def numero_es(valor, decimales=0):
    return f"{valor:,.{decimales}f}".replace(",", "_").replace(".", ",").replace("_", ".")


def main():
    carpeta = Path(__file__).resolve().parent
    entrada = carpeta / "output_costo" / "costo_mortalidad.xlsx"
    if not entrada.is_file():
        raise SystemExit("Falta output_costo/costo_mortalidad.xlsx. Ejecuta primero costo_mortalidad.py.")
    libro = load_workbook(entrada, read_only=True, data_only=True)
    try:
        if "por_centro" not in libro.sheetnames:
            raise SystemExit("El Excel no contiene la hoja por_centro.")
        filas = list(libro["por_centro"].values)
    finally:
        libro.close()
    columnas = list(filas[0])
    necesarias = {"centro_id", "nombre", "costo_kg_bruto_clp", "costo_kg_neto_clp"}
    if not necesarias.issubset(columnas):
        raise SystemExit("Faltan columnas de costo en la hoja por_centro.")
    registros = [dict(zip(columnas, fila)) for fila in filas[1:] if fila[0] is not None]
    # El registro con id 0 es una referencia agrupada, no un octavo centro.
    referencias = [fila for fila in registros if int(fila["centro_id"]) == 0]
    centros = [fila for fila in registros if int(fila["centro_id"]) != 0]
    if len(referencias) != 1 or not centros:
        raise SystemExit("Se necesitan centros y una referencia del resto (centro_id=0).")
    vistos = set()
    for fila in centros + referencias:
        cid = int(fila["centro_id"])
        costo = float(fila["costo_kg_bruto_clp"])
        if cid in vistos or not math.isfinite(costo) or costo <= 0:
            raise SystemExit("Hay centros duplicados o un costo bruto no válido.")
        vistos.add(cid)
    centros.sort(key=lambda fila: (-float(fila["costo_kg_bruto_clp"]), int(fila["centro_id"])))
    top = centros[0]
    referencia = referencias[0]
    costo_top = float(top["costo_kg_bruto_clp"])
    costo_resto = float(referencia["costo_kg_bruto_clp"])
    brecha = costo_top - costo_resto
    aumento_pct = (costo_top / costo_resto - 1) * 100
    salida = carpeta / "output_persona5"
    salida.mkdir(exist_ok=True)
    png = salida / "01_costo_alimento_por_centro.png"
    datos = salida / "01_datos_costos.csv"
    ficha = salida / "01_interpretacion_costos.md"

    plt.rcParams.update({"font.size": 12, "axes.spines.top": False,
                         "axes.spines.right": False})
    fig, ax = plt.subplots(figsize=(12, 7.8))
    posiciones = list(range(len(centros)))
    valores = [float(fila["costo_kg_bruto_clp"]) for fila in centros]
    colores = ["#27647A"] + ["#CCD2D8"] * (len(centros) - 1)
    ax.barh(posiciones, valores, color=colores, height=0.60)
    ax.set_yticks(posiciones, [fila["nombre"].removeprefix("Centro ") for fila in centros])
    ax.invert_yaxis()
    ax.get_yticklabels()[0].set_fontweight("bold")
    for y, valor in zip(posiciones, valores):
        ax.text(valor + costo_top * 0.015, y, "$" + numero_es(valor),
                va="center", fontsize=12, color="#374151")
    ax.axvline(costo_resto, color="#876129", linestyle="--", linewidth=1.5,
               label=f"Resto agrupado: ${numero_es(costo_resto)}/kg")
    ax.legend(loc="lower left", bbox_to_anchor=(0, 1.01), frameon=False,
              fontsize=11, borderaxespad=0)
    ax.set_xlim(0, costo_top * 1.18)
    ax.set_xlabel("Costo estimado de alimento (CLP por kg de crecimiento bruto)", labelpad=12)
    ax.xaxis.set_major_formatter(FuncFormatter(lambda valor, _: numero_es(valor)))
    ax.grid(axis="x", color="#E5E7EB", linewidth=0.7)
    ax.set_axisbelow(True)
    ax.spines["left"].set_visible(False)
    ax.tick_params(axis="y", length=0, pad=10)
    fig.text(0.17, 0.94, "Costo de alimento por centro (CLP/kg de crecimiento bruto)",
             fontsize=17, fontweight="bold", ha="left")
    mensaje = (f"{top['nombre']}: ${numero_es(costo_top)}/kg, "
               f"{numero_es(aumento_pct, 2)}% sobre el resto agrupado.")
    fig.text(0.17, 0.90, mensaje, fontsize=12, color="#374151")
    fig.subplots_adjust(left=0.17, right=0.98, top=0.80, bottom=0.22)
    fig.text(0.17, 0.06,
             "Fuente: SQLite operativo y CSV de precios; costo_mortalidad.py / costo_mortalidad.xlsx, hoja por_centro.\n"
             "Observaciones limpias de 2025. Precio mensual: promedio simple de tipos disponibles, aplicado a todos los centros.\n"
             "Referencia: costo total / crecimiento bruto total del resto. Incluye alimento; no es el costo operacional completo.",
             fontsize=9, color="#555555", va="bottom")
    fig.savefig(png, dpi=200, facecolor="white")
    plt.close(fig)

    # La tabla completa conserva ambas definiciones y la fila de referencia.
    with datos.open("w", encoding="utf-8-sig", newline="") as archivo:
        escritor = csv.DictWriter(archivo, fieldnames=columnas, lineterminator="\n")
        escritor.writeheader()
        escritor.writerows(centros + referencias)
    texto = f"""# Costo estimado de alimento por centro

**Pregunta:** ¿Qué centro tiene mayor costo de alimento por kg de crecimiento bruto?

**Origen del hallazgo:** medido bajo las reglas de valorización del alimento adoptadas por el equipo.

## Interpretación

{top['nombre']} tiene el mayor costo estimado de alimento por kg de crecimiento bruto: ${numero_es(costo_top, 2)}/kg. El resto agrupado presenta ${numero_es(costo_resto, 2)}/kg. La brecha es ${numero_es(brecha, 2)}/kg, equivalente a {numero_es(aumento_pct, 2)}% sobre la referencia.

La referencia agrupa costo y crecimiento del resto de los centros antes de dividir. No es el promedio simple de sus costos/kg. La figura utiliza crecimiento bruto, que es la definición usada por costo_mortalidad.py para seleccionar el centro de mayor costo. La tabla de datos conserva también el costo por crecimiento neto estimado.

## Fórmula y supuestos

Costo/kg bruto = suma(alimento_kg × precio mensual de referencia) / suma(crecimiento_kg).

La operación no identifica el tipo de alimento consumido. Se usa, por tanto, un precio mensual de referencia: promedio simple de los tipos con precio válido ese mes. Los tipos faltantes no se imputan, pero su ausencia puede modificar la composición de ese promedio. Se utiliza el mismo precio mensual en todos los centros. Se interpreta crecimiento_kg como bruto, según la ficha del FCR.

## Límites

El resultado incluye únicamente alimento. No contiene todos los costos operacionales ni demuestra por sí solo la causa del mayor costo. Los importes del gráfico se muestran redondeados a pesos; porcentajes y brechas se calculan a partir de los valores disponibles en el Excel, guardados con tres decimales.

## Reproducción

Con las salidas de limpieza actualizadas, ejecutar:

```bash
python costo_mortalidad.py
python graficos_costos.py
```

Entrada: output_costo/costo_mortalidad.xlsx, hoja por_centro. Salidas: output_persona5/01_costo_alimento_por_centro.png, 01_datos_costos.csv y 01_interpretacion_costos.md.

Uso de IA: ChatGPT ayudó a preparar y revisar la visualización y su interpretación.
"""
    ficha.write_text(texto, encoding="utf-8")
    print(f"OK costo de alimento: {len(centros)} centros")
    print(mensaje)
    print(f"Brecha: ${numero_es(brecha, 2)}/kg")
    print(f"Gráfico: {png}")
    print(f"Datos: {datos}")
    print(f"Interpretación: {ficha}")


if __name__ == "__main__":
    main()
