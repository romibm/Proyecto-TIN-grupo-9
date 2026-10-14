"""Compara mortalidad observada y esperada según la basal para un evento."""

import argparse
import csv
import datetime as dt
import math
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
from openpyxl import load_workbook


def numero_es(valor, decimales=0):
    return f"{valor:,.{decimales}f}".replace(",", "_").replace(".", ",").replace("_", ".")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--evento", default="EV-8001")
    args = parser.parse_args()
    carpeta = Path(__file__).resolve().parent
    entrada = carpeta / "output_costo" / "costo_mortalidad.xlsx"
    metadatos = carpeta / "output_persona2" / "eventos_limpios.csv"
    if not entrada.is_file() or not metadatos.is_file():
        raise SystemExit("Faltan costo_mortalidad.xlsx o eventos_limpios.csv. Ejecuta los pasos anteriores.")
    libro = load_workbook(entrada, read_only=True, data_only=True)
    try:
        if "eventos" not in libro.sheetnames:
            raise SystemExit("El Excel no contiene la hoja eventos.")
        filas = list(libro["eventos"].values)
    finally:
        libro.close()
    columnas = list(filas[0])
    registros = [dict(zip(columnas, fila)) for fila in filas[1:]]
    seleccion = [fila for fila in registros if fila.get("evento_id") == args.evento]
    with metadatos.open(encoding="utf-8-sig", newline="") as archivo:
        detalles = [fila for fila in csv.DictReader(archivo) if fila["evento_id"] == args.evento]
    if len(seleccion) != 1 or len(detalles) != 1:
        raise SystemExit("El evento debe existir una sola vez en cada fuente.")
    evento, detalle = seleccion[0], detalles[0]
    observadas = float(evento["muertes_observadas"])
    esperadas = float(evento["muertes_esperadas_basal"])
    exceso = float(evento["exceso_n"])
    if not all(math.isfinite(x) for x in (observadas, esperadas, exceso)) or min(observadas, esperadas) < 0:
        raise SystemExit("Hay valores de mortalidad inválidos.")
    # El Excel guarda estos resultados con un decimal; se admite ese redondeo.
    if not math.isclose(observadas - esperadas, exceso, rel_tol=0, abs_tol=0.11):
        raise SystemExit("El exceso del Excel no coincide con observadas menos esperadas.")
    inicio = dt.date.fromisoformat(detalle["fecha_inicio"])
    fin = dt.date.fromisoformat(detalle["fecha_fin"])
    if inicio > fin or evento["centro"] != detalle["centro"]:
        raise SystemExit("Los metadatos del evento no coinciden con el resultado.")
    jaulas = sorted(set(map(int, detalle["jaulas_afectadas"].split())))
    salida = carpeta / "output_persona5"
    salida.mkdir(exist_ok=True)
    png = salida / f"04_mortalidad_vs_basal_{args.evento}.png"
    datos = salida / f"04_datos_basal_{args.evento}.csv"
    ficha = salida / f"04_interpretacion_basal_{args.evento}.md"

    plt.rcParams.update({"font.size": 12, "axes.spines.top": False,
                         "axes.spines.right": False})
    fig, ax = plt.subplots(figsize=(11.5, 7.2))
    valores = [esperadas, observadas]
    ax.bar([0, 1], valores, width=0.52, color=["#CCD2D8", "#27647A"])
    etiquetas = [numero_es(esperadas, 1), numero_es(observadas)]
    for x, valor, etiqueta in zip([0, 1], valores, etiquetas):
        ax.text(x, valor + max(valores) * 0.025, etiqueta,
                ha="center", fontsize=15, fontweight="bold", color="#374151")
    ax.set_xticks([0, 1], ["Esperadas según basal\n(estimación)", "Observadas durante el evento\n(registros disponibles)"])
    ax.set_ylabel("Peces muertos en el intervalo (cantidad)")
    ax.set_ylim(0, max(valores) * 1.18 if max(valores) > 0 else 1)
    ax.yaxis.set_major_formatter(FuncFormatter(lambda valor, _: numero_es(valor)))
    ax.grid(axis="y", color="#E5E7EB", linewidth=0.7)
    ax.set_axisbelow(True)
    fig.text(0.12, 0.94, f"{evento['centro']}: mortalidad observada y basal (peces)",
             fontsize=17, fontweight="bold")
    mensaje = f"Exceso estimado sobre la basal: {numero_es(exceso, 1)} peces."
    fig.text(0.12, 0.89, mensaje, fontsize=13, color="#374151")
    fig.text(0.12, 0.84,
             f"{args.evento}: {evento['tipo']} | {inicio:%d/%m/%Y} a {fin:%d/%m/%Y} | "
             f"Jaulas {', '.join(map(str, jaulas))}", fontsize=10, color="#555555")
    fig.subplots_adjust(left=0.12, right=0.96, top=0.77, bottom=0.24)
    fig.text(0.12, 0.05,
             "Fuente: SQLite operativo y JSON sanitario; costo_mortalidad.py / costo_mortalidad.xlsx, hoja eventos.\n"
             "Basal por lote: muertes / suma de población en observaciones sin evento. Se aplica a la población en días de evento.\n"
             "Se comparan las mismas observaciones disponibles. El exceso es una estimación de referencia, no una atribución causal.",
             fontsize=9, color="#555555", va="bottom")
    fig.savefig(png, dpi=200, facecolor="white")
    plt.close(fig)

    fila_salida = dict(evento)
    fila_salida.update({"fecha_inicio": inicio.isoformat(), "fecha_fin": fin.isoformat(),
                        "jaulas_afectadas": " ".join(map(str, jaulas))})
    with datos.open("w", encoding="utf-8-sig", newline="") as archivo:
        escritor = csv.DictWriter(archivo, fieldnames=list(fila_salida), lineterminator="\n")
        escritor.writeheader()
        escritor.writerow(fila_salida)
    texto = f"""# Mortalidad observada frente a basal — {args.evento}

**Pregunta:** ¿Cuántas muertes se observaron en las jaulas del evento frente a las esperadas bajo la referencia basal?

## Interpretación

En {evento['centro']}, durante {args.evento} ({evento['tipo']}), se observaron {numero_es(observadas)} muertes en las jaulas {', '.join(map(str, jaulas))}, entre {inicio.isoformat()} y {fin.isoformat()}, ambos días incluidos. Aplicando la basal del equipo a las mismas observaciones se esperan {numero_es(esperadas, 1)} muertes. El exceso estimado es {numero_es(exceso, 1)} peces.

El observado es un conteo de registros conservados. El esperado y el exceso dependen de la referencia escogida. No se demuestra que todas las muertes del exceso fueran causadas por el evento.

## Fórmulas de persona 4

- Tasa basal por lote = suma de mortalidad_n / suma de poblacion en observaciones sin evento sanitario registrado.
- Muertes esperadas durante el evento = suma de tasa_basal_del_lote × poblacion de cada observación del evento.
- Exceso = muertes observadas - muertes esperadas.

La suma de población corresponde a exposición acumulada entre observaciones diarias, no a peces únicos. La tasa basal se aplica por lote; no se promedian tasas diarias con el mismo peso.

## Alcance y límites

Una fila de la tabla utilizada es un evento sanitario y sus observaciones de jaulas afectadas. El contraste utiliza las observaciones disponibles en la base limpia. No se imputa mortalidad para días o jaulas sin registros. Las filas excluidas por limpieza pueden cambiar los resultados.

Se supone que la tasa de los días sin evento representa la mortalidad esperable para ese lote durante el intervalo. La referencia incluye días previos y posteriores, por lo que describe una comparación retrospectiva. La ausencia de un evento registrado no garantiza ausencia de enfermedad.

El gráfico contiene dos cantidades comparables; el exceso se informa como su diferencia y no como una tercera cantidad que deba sumarse. El Excel conserva la tabla de eventos con un decimal y el CSV copia esa precisión.

## Reproducción

Con las salidas de limpieza disponibles:

```bash
python costo_mortalidad.py
python graficos_basal.py --evento {args.evento}
```

Entrada: output_costo/costo_mortalidad.xlsx, hoja eventos. Fechas y jaulas: output_persona2/eventos_limpios.csv. Salidas: output_persona5/04_mortalidad_vs_basal_{args.evento}.png, 04_datos_basal_{args.evento}.csv y 04_interpretacion_basal_{args.evento}.md.

Uso de IA: ChatGPT ayudó a preparar y revisar la visualización y su interpretación.
"""
    ficha.write_text(texto, encoding="utf-8")
    print(f"OK mortalidad frente a basal: {args.evento}; {evento['centro']}")
    print(f"Muertes observadas: {numero_es(observadas)}")
    print(f"Muertes esperadas según basal: {numero_es(esperadas, 1)}")
    print(mensaje)
    print(f"Gráfico: {png}")
    print(f"Datos: {datos}")
    print(f"Interpretación: {ficha}")


if __name__ == "__main__":
    main()
